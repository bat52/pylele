#!/usr/bin/env python3

"""
    Pylele Tests
"""

import unittest
import os
import csv

from json_tricks import load
from prettytable import from_csv

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

import sys
sys.path.append(os.path.join(os.path.dirname(__file__), "../"))

from pylele.api.core import test_api, DEFAULT_TEST_DIR
from pylele.api.utils import make_or_exist_path

REPORT_COLS=["subdir_level_1",
             "subdir_level_2", 
             "subdir_level_3",
             "subdir_level_4",
             "filename",
             "stl_file_size",
             "pass",
             "render_time",
             "volume",
             "convex_hull_volume",
             "bounding_box_x",
             "bounding_box_y",
             "bounding_box_z",
             "datetime",
             ]
REPORT_COLS_COMPACT=[
             "subdir_level_1",
             "subdir_level_2", 
             "subdir_level_3",
             "subdir_level_4",
             "pass",
             "volume",
             ]
REPORT_EXCLUDE={"subdir_level_3":"mock"}

def json_to_csv(directory, output_csv, include_filename=True,
                filter_out=REPORT_EXCLUDE,
                column_order=REPORT_COLS):
    """
    Recursively searches for JSON files in the specified directory and its subdirectories,
    extracts data from files with '_rpt.json' in their filename using json-tricks,
    and saves it to a CSV file with separate columns for each subdirectory level.
    Optionally filters out rows based on specified criteria and orders columns based on a custom list.

    Parameters:
    - directory (str): Path to the root directory to search for JSON files.
    - output_csv (str): Path to the output CSV file.
    - include_filename (bool): If True, includes the filename as a column in the CSV.
    - filter_out (dict): Dictionary of column-value pairs to filter out from the output.
    - column_order (list): List of column names specifying the desired order in the output.
    """
    rows = []
    headers = set()
    filter_out = filter_out or {}

    # Walk through the directory structure
    for root, _, files in os.walk(directory):
        for file in files:
            if file.endswith('_rpt.json'):
                file_path = os.path.join(root, file)
                try:
                    # Read JSON data from file using json-tricks
                    with open(file_path, 'r') as f:
                        data = load(f)  # Load JSON with json-tricks
                    
                    # Extract subdirectory path relative to root directory
                    subdirectory_path = os.path.relpath(root, directory)
                    subdirectories = subdirectory_path.split(os.sep)
                    
                    # Flatten JSON data to a single row dictionary
                    row = {}
                    for key, value in data.items():
                        row[key] = value
                        headers.add(key)  # Keep track of all column headers

                    # Add columns for each directory level
                    for i, subdir in enumerate(subdirectories):
                        col_name = f"subdir_level_{i + 1}"
                        row[col_name] = subdir
                        headers.add(col_name)
                    
                    # Optionally add filename as a column
                    if include_filename:
                        row["filename"] = file
                        headers.add("filename")

                    # Filter out rows based on filter_out criteria
                    if any(row.get(k) == v for k, v in filter_out.items()):
                        continue  # Skip this row if it matches filter criteria
                    
                    rows.append(row)
                
                except (ValueError, IOError) as e:
                    print(f"Error reading {file_path}: {e}")

    # Determine the final order of headers based on column_order
    if column_order:
        # Only include columns specified in column_order and present in headers
        ordered_headers = [col for col in column_order if col in headers]
    else:
        # If no column_order specified, include all headers in alphabetical order
        ordered_headers = sorted(headers)

    # Write to CSV file
    with open(output_csv, 'w', newline='') as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=ordered_headers, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)

    print(f"Data saved to {output_csv}")

def print_csv(fname):
    """ Print Table from .csv file """
    with open(fname) as fp:
        table = from_csv(fp,delimiter=',')
        print(table.get_string(sortby="pass",fields=REPORT_COLS_COMPACT))

def csv_to_xls(csv_file, xls_file):
    """
    Converts a CSV file into an Excel .xlsx file with a table that has default filtering active.
    
    Parameters:
    - csv_file (str): Path to the input CSV file.
    - xls_file (str): Path to the output Excel file.
    """
    # Create a new workbook and active worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    # Open the CSV file and read its content
    with open(csv_file, newline='') as f:
        reader = csv.reader(f)
        headers = next(reader)  # Get the headers

        # Write headers
        ws.append(headers)
        
        # Write the rest of the rows
        for row in reader:
            new_row = []
            for cell in row:
                # Try to convert each cell to a number
                try:
                    # Convert to float and then to int if no decimal part
                    numeric_value = float(cell)
                    cell_value = int(numeric_value) if numeric_value.is_integer() else numeric_value
                except ValueError:
                    # If conversion fails, keep it as a string
                    cell_value = cell
                new_row.append(cell_value)
            ws.append(new_row)

    # Define the range for the table (including all rows and columns)
    last_column = chr(64 + len(headers))  # ASCII to column name (e.g., 'A', 'B', ..., 'Z')
    table_ref = f"A1:{last_column}{ws.max_row}"

    # Create a table with autofiltering enabled
    table = Table(displayName="TableWithFilter", ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=True
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    # Save the workbook as an .xlsx file
    wb.save(xls_file)

    print(f"Data from '{csv_file}' has been saved to '{xls_file}' with default filtering enabled.")

class PyleleTestMethods(unittest.TestCase):
    """Pylele Test Class"""

    ## API
    def test_mock_api(self):
        """Test Mock API"""
        test_api(api="mock")

    def test_cadquery_api(self):
        """Test Cadquery API"""
        test_api(api="cadquery")

    def test_blender_api(self):
        """Test Blender API"""
        test_api(api="blender")

    def test_trimesh_api(self):
        """Test Trimesh API"""
        test_api(api="trimesh")

    def test_solid2_api(self):
        """Test SolidPython2 API"""
        test_api(api="solid2")

    def test_manifold_api(self):
        """Test Manifold API"""
        test_api(api="manifold")

    ## Solid Parts
    from pylele.parts.tube import test_tube, test_tube_mock
    from pylele.parts.screw import test_screw, test_screw_mock
    from pylele.parts.import3d import test_import3d
    from pylele.parts.scad_example import test_scad_example
    from pylele.parts.rounded_box import test_rounded_box, test_rounded_box_mock
    from pylele.parts.tunable_saddle import test_tunable_saddle, test_tunable_saddle_mock
    from pylele.parts.tunable_bridge import test_tunable_bridge, test_tunable_bridge_mock

    ## Pylele Individual Parts
    from pylele.pylele2.frets import test_frets, test_frets_mock
    from pylele.pylele2.fretboard import test_fretboard, test_fretboard_mock
    from pylele.pylele2.fretboard_dots import (
        test_fretboard_dots,
        test_fretboard_dots_mock,
    )
    from pylele.pylele2.fretboard_spines import (
        test_fretboard_spines,
        test_fretboard_spines_mock,
    )
    from pylele.pylele2.fretboard_joint import (
        test_fretboard_joint,
        test_fretboard_joint_mock,
    )
    from pylele.pylele2.top import test_top, test_top_mock
    from pylele.pylele2.strings import test_strings, test_strings_mock
    from pylele.pylele2.nut import test_nut, test_nut_mock
    from pylele.pylele2.spines import test_spines, test_spines_mock
    from pylele.pylele2.head import test_head, test_head_mock
    from pylele.pylele2.neck_joint import test_neck_joint, test_neck_joint_mock
    from pylele.pylele2.neck import test_neck, test_neck_mock
    from pylele.pylele2.bridge import test_bridge, test_bridge_mock
    from pylele.pylele2.guide import test_guide, test_guide_mock
    from pylele.pylele2.chamber import test_chamber, test_chamber_mock
    from pylele.pylele2.peg import test_peg, test_peg_mock
    from pylele.pylele2.worm import test_worm, test_worm_mock
    from pylele.pylele2.tuners import test_tuners, test_tuners_mock
    from pylele.pylele2.body import test_body, test_body_mock
    from pylele.pylele2.texts import test_texts, test_texts_mock
    from pylele.pylele2.rim import test_rim, test_rim_mock
    from pylele.pylele2.worm_key import test_worm_key, test_worm_key_mock
    from pylele.pylele2.tail import test_tail, test_tail_mock
    from pylele.pylele2.brace import test_brace, test_brace_mock
    from pylele.pylele2.soundhole import test_soundhole, test_soundhole_mock
    from pylele.pylele2.neck_bend import test_neck_bend, test_neck_bend_mock

    ## Assemblies
    from pylele.pylele2.fretboard_assembly import (
        test_fretboard_assembly,
        test_fretboard_assembly_mock,
    )
    from pylele.pylele2.neck_assembly import test_neck_assembly, test_neck_assembly_mock
    from pylele.pylele2.top_assembly import test_top_assembly, test_top_assembly_mock
    from pylele.pylele2.bottom_assembly import (
        test_bottom_assembly,
        test_bottom_assembly_mock,
    )
    from pylele.pylele2.all_assembly import test_all_assembly, test_all_assembly_mock

    def test_zz_report(self):
        """ Generate Test Report """
        print("# Generate .xlsx Test Report")

        basefname = os.path.join(DEFAULT_TEST_DIR,"test_report")
        csvfname = basefname + '.csv'
        xlsfname = basefname + '.xlsx'
        json_to_csv(DEFAULT_TEST_DIR, csvfname)
        print_csv(csvfname)
        csv_to_xls(csvfname, xlsfname)

def test_main():
    """ Launch all tests """      
    
    if os.getenv("GITHUB_ACTIONS") == "true":
        print("Running in GitHub Actions: simplified output to log file")

        # text test report
        make_or_exist_path(DEFAULT_TEST_DIR)
        with open(os.path.join(DEFAULT_TEST_DIR,"test_log.txt"), "w") as f:
            runner = unittest.TextTestRunner(stream=f, verbosity=3)
            unittest.main(testRunner=runner)

    else:
        print("Running Locally: all output to shell")
        unittest.main()

if __name__ == "__main__":
    test_main()
